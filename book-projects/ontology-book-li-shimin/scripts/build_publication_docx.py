"""Build editable Word documents from canonical reader text, without rewriting prose."""
from pathlib import Path
import re, json, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'book_project/05_exports'

def reader_text(p):
    text = p.read_text(encoding='utf-8')
    if 'chapter_' in p.name:
        title = text.splitlines()[0]
        body = text.split('<!-- BODY_START -->')[1].split('<!-- BODY_END -->')[0].strip()
        tail = text.split('<!-- BODY_END -->')[1].split('## 編輯紀錄（非出版正文）')[0].strip()
        return title+'\n\n'+body+'\n\n'+tail+'\n'
    return text

def font(style, size, face='Microsoft JhengHei'):
    style.font.name = face
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), face)

def inline(par, text):
    # Links remain real hyperlinks; inline emphasis does not leak Markdown syntax.
    pattern = r'\[([^\]]+)\]\((https?://[^\s)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`'
    pos = 0
    for m in re.finditer(pattern, text):
        par.add_run(text[pos:m.start()])
        if m.group(1):
            rel = par.part.relate_to(m.group(2), RT.HYPERLINK, is_external=True)
            link = OxmlElement('w:hyperlink'); link.set(qn('r:id'),rel)
            run=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
            color=OxmlElement('w:color'); color.set(qn('w:val'),'245D74'); prop.append(color)
            run.append(prop); t=OxmlElement('w:t'); t.text=m.group(1); run.append(t); link.append(run); par._p.append(link)
        else:
            run=par.add_run(m.group(3) or m.group(4)); run.bold=bool(m.group(3))
        pos=m.end()
    par.add_run(text[pos:])

def build(text, path, full=False):
    doc=Document(); sec=doc.sections[0]
    sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.1); sec.bottom_margin=Cm(2); sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.2)
    sec.header_distance=Cm(.8); sec.footer_distance=Cm(.8)
    font(doc.styles['Normal'],11)
    normal=doc.styles['Normal'].paragraph_format; normal.line_spacing=1.35; normal.space_after=Pt(7); normal.widow_control=True
    for name,size in [('Title',24),('Subtitle',13),('Heading 1',17),('Heading 2',14),('Heading 3',12),('List Bullet',11),('List Number',11)]:
        font(doc.styles[name],size)
    for name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        doc.styles[name].paragraph_format.keep_with_next=True
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))): border.getparent().remove(border)
    doc.styles['Subtitle'].font.italic=False
    if not full: doc.styles['Title'].font.size=Pt(17)
    doc.styles['Heading 1'].paragraph_format.page_break_before=True
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(14)
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); run._r.addnext(fld)
    doc.core_properties.author='李詩民'; doc.core_properties.title=text.splitlines()[0].lstrip('# ')
    lines=text.splitlines(); i=0; first=True; in_code=False; in_toc=False; expected=[]
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line or line=='---' or line.startswith('<!--'): continue
        if line.startswith('```'): in_code=not in_code; continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'): rows.append(lines[i].strip()); i+=1
            vals=[[c.strip() for c in r.strip('|').split('|')] for r in rows if not re.fullmatch(r'[| :\-]+',r)]
            table=doc.add_table(rows=len(vals), cols=max(map(len,vals))); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
            n=len(vals[0]); widths=([2.1,4,10.5] if n==3 else [16.6/n]*n)
            if n==4: widths=[2.2,3.7,5.4,5.3]
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            table._tbl.tblPr.append(borders)
            for r,row in enumerate(vals):
                trpr=table.rows[r]._tr.get_or_add_trPr(); nosplit=OxmlElement('w:cantSplit'); trpr.append(nosplit)
                if r==0: trpr.append(OxmlElement('w:tblHeader'))
                for c,val in enumerate(row):
                    cell=table.cell(r,c); cell.width=Cm(widths[c]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    pr=cell._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
                    for edge in ['top','left','bottom','right']:
                        el=OxmlElement('w:'+edge); el.set(qn('w:w'),'95'); el.set(qn('w:type'),'dxa'); margins.append(el)
                    pr.append(margins)
                    fill=OxmlElement('w:shd'); fill.set(qn('w:fill'),'E7EEF2' if r==0 else 'FFFFFF'); pr.append(fill)
                    p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.2; inline(p,val); expected.append(val)
                    for run in p.runs: run.font.size=Pt(9.5); run.bold=r==0
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        heading=re.match(r'^(#{1,6}) (.+)$',line)
        if heading:
            level=len(heading.group(1)); value=heading.group(2)
            if first: p=doc.add_paragraph(style='Title'); first=False
            elif full and value=='目錄':
                p=doc.add_paragraph('目錄'); p.paragraph_format.page_break_before=True
                p.runs[0].bold=True; p.runs[0].font.size=Pt(17)
                fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u'); doc.add_paragraph()._p.append(fld); in_toc=True; continue
            elif full and not any(p.style.name=='Heading 1' for p in doc.paragraphs) and level==2: p=doc.add_paragraph(style='Subtitle')
            else:
                in_toc=False; p=doc.add_paragraph(style='Heading '+str(min(level,3)))
            inline(p,value); expected.append(value); continue
        if in_toc and line.startswith('- '): continue
        if line.startswith('- '): p=doc.add_paragraph(style='List Bullet'); line=line[2:]
        elif re.match(r'^\d+\. ',line):
            p=doc.add_paragraph(); p.paragraph_format.left_indent=Cm(.55); p.paragraph_format.first_line_indent=Cm(-.55)
        elif line.startswith('> '): p=doc.add_paragraph(style='Quote'); line=line[2:]
        else: p=doc.add_paragraph()
        inline(p,line); expected.append(line)
    if not full and len(doc.paragraphs)>2:
        doc.paragraphs[-2].paragraph_format.keep_with_next=True
    if full:
        paragraphs=doc.paragraphs
        for idx,p in enumerate(paragraphs):
            if p.style.name=='Heading 1' and idx>=2 and paragraphs[idx-1].style.name=='List Bullet':
                paragraphs[idx-2].paragraph_format.keep_with_next=True
    settings=doc.settings.element; update=OxmlElement('w:updateFields'); update.set(qn('w:val'),'true'); settings.append(update)
    path.parent.mkdir(parents=True,exist_ok=True); doc.save(path)
    return {'path':path.relative_to(ROOT).as_posix(),'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'source_blocks':len(expected),'render_status':'pending'}

def main():
    OUT.mkdir(parents=True,exist_ok=True); entries=[]
    for n in range(1,17):
        src=ROOT/f'book_project/02_chapters/chapter_{n:02}/chapter_{n:02}_main.md'; text=reader_text(src)
        md=OUT/f'reader_chapters/chapter_{n:02}_main.md'; md.parent.mkdir(parents=True,exist_ok=True); md.write_text(text,encoding='utf-8',newline='\n')
        e=build(text,OUT/f'docx/chapter_{n:02}_main.docx'); e.update(source=src.relative_to(ROOT).as_posix(),source_sha256=hashlib.sha256(src.read_bytes()).hexdigest()); entries.append(e)
    src=ROOT/'book_project/01_master/full_book.md'; e=build(reader_text(src),OUT/'docx/full_book.docx',True); e.update(source=src.relative_to(ROOT).as_posix(),source_sha256=hashlib.sha256(src.read_bytes()).hexdigest()); entries.append(e)
    (OUT/'docx_manifest.json').write_text(json.dumps({'count':len(entries),'entries':entries},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'created':len(entries),'output':str(OUT),'render_status':'pending'},ensure_ascii=False))

if __name__=='__main__': main()
